import json
from datetime import date, datetime

import pytest
from openpyxl import load_workbook

from src import excel_log
from src.analyzer import AnalysisFailed
from src.config import Config
from src.discord_listener import image_extension
from src.excel_log import COLUMNS, DayCounts, ExcelLog
from src.worker import MessageGroup, QueueItem, Worker

NOW = datetime(2026, 10, 22, 14, 5, 33)

CAR = {"is_vehicle": True, "vehicle_type": "car", "brand": "Toyota", "model": "Corolla",
       "estimated_year": 2015, "body_type": "sedan", "confidence": 85, "description": "silver sedan"}
BUS = {"is_vehicle": False, "vehicle_type": None, "estimated_year": None, "confidence": 90,
       "description": "bus"}


def claude_says(answer):
    text = answer if isinstance(answer, str) else json.dumps(answer)
    return json.dumps({"is_error": False, "result": text})


class FakeMessage:
    def __init__(self):
        self.replies, self.removed = [], []

    async def reply(self, text, **options):
        self.replies.append((text, options))

    async def remove_reaction(self, emoji, user):
        self.removed.append(emoji)


class FakeAttachment:
    def __init__(self, data=b"photo", content_type="image/jpeg", filename="a.jpg", fail=False):
        self.data, self.content_type, self.filename, self.fail = data, content_type, filename, fail

    async def read(self):
        if self.fail:
            raise OSError("download broke")
        return self.data


@pytest.fixture
def env(tmp_path, monkeypatch):
    monkeypatch.setattr(excel_log, "SECONDS_BETWEEN_TRIES", 0)
    config = Config("fake-token", 1, "model", tmp_path)
    excel = ExcelLog(tmp_path)
    counts = DayCounts(day=date(2026, 10, 22))
    return config, excel, counts


def make_worker(env, answers):
    """A worker whose 'claude -p' returns the given answers one after another (an Exception is raised)."""
    config, excel, counts = env
    calls = []

    async def fake_analyze(cfg, image_bytes, extension, uploader_text):
        calls.append(uploader_text)
        answer = answers[min(len(calls), len(answers)) - 1]
        if isinstance(answer, Exception):
            raise answer
        return answer

    worker = Worker(config, excel, counts, bot_user="bot", analyze=fake_analyze)
    worker.calls = calls
    return worker


def make_item(message=None, attachment=None, remaining=1, note="", name="Calvin"):
    group = MessageGroup(message=message or FakeMessage(), remaining=remaining)
    return QueueItem(group=group, attachment=attachment or FakeAttachment(), extension=".jpg",
                     uploader_name=name, uploader_id=123456789012345678, note=note,
                     message_link="https://discord.com/channels/1/2/3", received_at=NOW)


def log_rows(env):
    sheet = load_workbook(env[1].path)["Log"]
    return [dict(zip(COLUMNS, [c.value for c in row])) for row in sheet.iter_rows(min_row=2)]


async def test_vehicle_photo_end_to_end(env):
    worker = make_worker(env, [claude_says(CAR)])
    item = make_item(note="my car!")
    await worker.process(item)

    text, options = item.group.message.replies[0]
    assert text.startswith("🚗 Car: Toyota Corolla, sedan, about 2015 (85% sure)")
    assert text.endswith("Today: 1 vehicle (1 car, 0 motorcycles)")
    assert options["mention_author"] is False
    assert item.group.message.removed == ["⏳"]

    row = log_rows(env)[0]
    assert (row["Status"], row["Brand"], row["Note"]) == ("OK", "Toyota", "my car!")
    assert (row["Today's total"], row["Today's cars"], row["Today's motorcycles"]) == (1, 1, 0)
    assert row["Uploader ID"] == "123456789012345678"
    assert row["Photo file"].endswith("2026-10-22 14-05-33 Toyota-Corolla.jpg")
    assert row["Discord message link"] == "https://discord.com/channels/1/2/3"
    assert (env[0].photos_dir / "2026-10-22 14-05-33 Toyota-Corolla.jpg").exists()


async def test_not_a_vehicle_is_replied_logged_saved_and_not_counted(env):
    worker = make_worker(env, [claude_says(BUS)])
    item = make_item()
    await worker.process(item)

    assert "(I see: bus)" in item.group.message.replies[0][0]
    row = log_rows(env)[0]
    assert (row["Status"], row["Brand"], row["Description"]) == ("Not a vehicle", "Not a vehicle", "bus")
    assert row["Today's total"] == 0
    assert any("not-a-vehicle" in p.name for p in env[0].photos_dir.iterdir())


async def test_low_confidence_is_flagged_but_still_counted(env):
    worker = make_worker(env, [claude_says({**CAR, "confidence": 30})])
    item = make_item()
    await worker.process(item)
    assert "I'm not sure" in item.group.message.replies[0][0]
    row = log_rows(env)[0]
    assert (row["Low confidence"], row["Today's total"]) == ("Yes", 1)


@pytest.mark.parametrize(
    "answer",
    [
        claude_says("this is not json"),
        claude_says({**CAR, "confidence": "high"}),
        claude_says({**CAR, "estimated_year": "unknown"}),
        AnalysisFailed("claude -p took longer than 120 seconds"),
        RuntimeError("something unexpected"),
    ],
)
async def test_bad_answers_give_the_failure_reply_an_error_row_and_no_crash(env, answer):
    worker = make_worker(env, [answer])
    item = make_item()
    await worker.process(item)

    assert item.group.message.replies[0][0] == "Recognition failed, please send the photo again"
    row = log_rows(env)[0]
    assert row["Status"] == "Error" and row["Today's total"] == 0 and row["Description"]
    assert worker.calls and len(worker.calls) == 2  # one retry, no more
    assert any("recognition-failed" in p.name for p in env[0].photos_dir.iterdir())


async def test_a_second_try_can_rescue_a_failed_first_try(env):
    worker = make_worker(env, [claude_says("garbage"), claude_says(CAR)])
    item = make_item()
    await worker.process(item)
    assert item.group.message.replies[0][0].startswith("🚗")
    assert len(worker.calls) == 2


async def test_failed_download_is_an_error_row_without_a_photo(env):
    worker = make_worker(env, [claude_says(CAR)])
    item = make_item(attachment=FakeAttachment(fail=True))
    await worker.process(item)
    assert item.group.message.replies[0][0].startswith("Recognition failed")
    assert log_rows(env)[0]["Photo file"] in ("", None)
    assert worker.calls == []


async def test_the_hourglass_comes_off_only_after_the_last_image_of_a_message(env):
    worker = make_worker(env, [claude_says(CAR)])
    message = FakeMessage()
    first = make_item(message, remaining=2)
    second = make_item(message, remaining=2)
    first.group = second.group = MessageGroup(message=message, remaining=2)

    await worker.process(first)
    assert message.removed == []
    await worker.process(second)
    assert message.removed == ["⏳"] and len(message.replies) == 2
    assert [r["Today's total"] for r in log_rows(env)] == [1, 2]


async def test_the_queue_works_through_photos_one_at_a_time_in_order(env):
    worker = make_worker(env, [claude_says(CAR)])
    import asyncio
    running, peak = 0, 0
    original = worker.analyze

    async def watching(*args):
        nonlocal running, peak
        running += 1
        peak = max(peak, running)
        await asyncio.sleep(0.01)
        try:
            return await original(*args)
        finally:
            running -= 1

    worker.analyze = watching
    items = [make_item(name=f"user{n}") for n in range(4)]
    for item in items:
        worker.submit(item)
    task = asyncio.create_task(worker.run_forever())
    await worker.queue.join()
    task.cancel()

    assert peak == 1
    assert [r["Uploader name"] for r in log_rows(env)] == ["user0", "user1", "user2", "user3"]
    assert [r["Today's total"] for r in log_rows(env)] == [1, 2, 3, 4]


async def test_a_crash_inside_one_photo_does_not_stop_the_queue(env, monkeypatch, capsys):
    import asyncio
    worker = make_worker(env, [claude_says(CAR)])
    bad, good = make_item(), make_item()
    monkeypatch.setattr(worker, "_download", None)  # makes process() blow up for the first one
    worker.submit(bad)
    task = asyncio.create_task(worker.run_forever())
    await worker.queue.join()
    monkeypatch.undo()
    worker.submit(good)
    await worker.queue.join()
    task.cancel()
    assert "ERROR while handling a photo" in capsys.readouterr().out
    assert good.group.message.replies


async def test_the_count_starts_again_on_a_new_date(env):
    worker = make_worker(env, [claude_says(CAR)])
    await worker.process(make_item())
    next_day = make_item()
    next_day.received_at = datetime(2026, 10, 23, 0, 0, 5)
    await worker.process(next_day)
    assert [r["Today's total"] for r in log_rows(env)] == [1, 1]


@pytest.mark.parametrize(
    "content_type, filename, expected",
    [
        ("image/jpeg", "a.jpg", ".jpg"),
        ("image/png", "a.png", ".png"),
        ("image/webp", "x", ".webp"),
        (None, "photo.JPEG", ".jpg"),
        ("application/pdf", "a.pdf", None),
        ("video/mp4", "a.mp4", None),
        ("image/svg+xml", "a.svg", None),
        (None, "notes.txt", None),
    ],
)
def test_which_attachments_count_as_photos(content_type, filename, expected):
    assert image_extension(FakeAttachment(content_type=content_type, filename=filename)) == expected
