import zipfile
from datetime import date, datetime

import pytest
from openpyxl import load_workbook

from src import excel_log
from src.excel_log import COLUMNS, DayCounts, ExcelLog

NOW = datetime(2026, 10, 22, 14, 5, 33)


@pytest.fixture(autouse=True)
def no_waiting(monkeypatch):
    monkeypatch.setattr(excel_log, "SECONDS_BETWEEN_TRIES", 0)


def make_row(**values):
    row = {"Time": NOW, "Uploader name": "Calvin", "Uploader ID": "123456789012345678", "Status": "OK",
           "Vehicle type": "Car", "Brand": "Toyota", "Model": "Corolla", "Estimated year": 2015,
           "Confidence": 85, "Today's total": 1, "Today's cars": 1, "Today's motorcycles": 0}
    row.update(values)
    return row


def read_back(path):
    sheet = load_workbook(path)["Log"]
    return [[cell.value for cell in row] for row in sheet.iter_rows()], sheet


def test_creates_the_file_with_header_and_adds_rows_at_the_bottom(tmp_path):
    log = ExcelLog(tmp_path)
    assert log.append_row_blocking(make_row(Brand="First"))
    assert log.append_row_blocking(make_row(Brand="Second"))
    rows, _ = read_back(log.path)
    assert rows[0] == COLUMNS
    assert [r[COLUMNS.index("Brand")] for r in rows[1:]] == ["First", "Second"]


def test_existing_rows_are_never_changed(tmp_path):
    log = ExcelLog(tmp_path)
    log.append_row_blocking(make_row(Brand="Old"))
    before, _ = read_back(log.path)
    log.append_row_blocking(make_row(Brand="New"))
    after, _ = read_back(log.path)
    assert after[:2] == before


def test_user_id_is_kept_as_exact_text_and_time_as_a_date(tmp_path):
    log = ExcelLog(tmp_path)
    log.append_row_blocking(make_row())
    rows, _ = read_back(log.path)
    assert rows[1][COLUMNS.index("Uploader ID")] == "123456789012345678"
    assert rows[1][0] == NOW


def test_no_formulas_anywhere(tmp_path):
    log = ExcelLog(tmp_path)
    log.append_row_blocking(make_row())
    with zipfile.ZipFile(log.path) as archive:
        assert b"<f>" not in archive.read("xl/worksheets/sheet1.xml")


def test_locked_file_keeps_rows_in_backup_then_adds_them_in_order(tmp_path, monkeypatch, capsys):
    log = ExcelLog(tmp_path)
    log.append_row_blocking(make_row(Brand="A"))

    real_replace = excel_log.os.replace
    def locked(*args):
        raise PermissionError("file is open in Excel")
    monkeypatch.setattr(excel_log.os, "replace", locked)

    assert log.append_row_blocking(make_row(Brand="B")) is False
    assert log.append_row_blocking(make_row(Brand="C")) is False
    printed = capsys.readouterr().out
    assert "Excel file is open" in printed and "retrying" in printed
    assert log.pending_path.exists()
    assert len(log.pending_path.read_text(encoding="utf-8").splitlines()) == 2

    monkeypatch.setattr(excel_log.os, "replace", real_replace)
    assert log.append_row_blocking(make_row(Brand="D"))
    rows, _ = read_back(log.path)
    assert [r[COLUMNS.index("Brand")] for r in rows[1:]] == ["A", "B", "C", "D"]
    assert not log.pending_path.exists()


def test_waiting_rows_are_added_when_the_bot_restarts_and_a_row_is_saved(tmp_path, monkeypatch):
    log = ExcelLog(tmp_path)
    monkeypatch.setattr(excel_log.os, "replace", lambda *a: (_ for _ in ()).throw(PermissionError()))
    log.append_row_blocking(make_row(Brand="Waiting"))
    monkeypatch.undo()
    monkeypatch.setattr(excel_log, "SECONDS_BETWEEN_TRIES", 0)

    restarted = ExcelLog(tmp_path)
    restarted.append_row_blocking(make_row(Brand="Fresh"))
    rows, _ = read_back(restarted.path)
    assert [r[COLUMNS.index("Brand")] for r in rows[1:]] == ["Waiting", "Fresh"]


def test_nothing_crashes_even_if_the_backup_cannot_be_written(tmp_path, monkeypatch, capsys):
    log = ExcelLog(tmp_path)
    monkeypatch.setattr(excel_log.os, "replace", lambda *a: (_ for _ in ()).throw(PermissionError()))
    log.pending_path.mkdir()  # a folder where the backup file should be
    assert log.append_row_blocking(make_row()) is False
    assert "could not save the row" in capsys.readouterr().out


def test_counts_are_rebuilt_from_todays_ok_rows_only(tmp_path):
    log = ExcelLog(tmp_path)
    yesterday = datetime(2026, 10, 21, 23, 0, 0)
    log.append_row_blocking(make_row(Time=yesterday))
    log.append_row_blocking(make_row())
    log.append_row_blocking(make_row(**{"Vehicle type": "Motorcycle"}))
    log.append_row_blocking(make_row(Status="Not a vehicle", **{"Vehicle type": ""}))
    log.append_row_blocking(make_row(Status="Error", **{"Vehicle type": ""}))
    counts = log.rebuild_counts(date(2026, 10, 22))
    assert (counts.total, counts.cars, counts.motorcycles) == (2, 1, 1)


def test_counts_include_rows_waiting_in_the_backup_file(tmp_path, monkeypatch):
    log = ExcelLog(tmp_path)
    monkeypatch.setattr(excel_log.os, "replace", lambda *a: (_ for _ in ()).throw(PermissionError()))
    log.append_row_blocking(make_row())
    counts = ExcelLog(tmp_path).rebuild_counts(date(2026, 10, 22))
    assert counts.total == 1


def test_day_counts_restart_on_a_new_date():
    counts = DayCounts(day=date(2026, 10, 22))
    counts.add("car")
    counts.add("Motorcycle")
    counts.add("bus")  # ignored
    assert (counts.total, counts.cars, counts.motorcycles) == (2, 1, 1)
    counts.roll_to(date(2026, 10, 23))
    assert (counts.total, counts.cars, counts.motorcycles) == (0, 0, 0)
