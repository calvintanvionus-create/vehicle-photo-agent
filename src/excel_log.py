"""The Excel log: append-only, no formulas, today's counts, retry and backup (section 8)."""

import asyncio
import json
import os
import threading
import time
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from .validator import STATUS_OK

SHEET = "Log"
LOG_NAME = "vehicle_log.xlsx"
PENDING_NAME = "pending_rows.jsonl"
SAVE_TRIES = 3
SECONDS_BETWEEN_TRIES = 2.0
TIME_FORMAT = "yyyy-mm-dd hh:mm:ss"

COLUMNS = [
    "Time",
    "Uploader name",
    "Uploader ID",
    "Status",
    "Vehicle type",
    "Brand",
    "Model",
    "Estimated year",
    "Body type",
    "Confidence",
    "Low confidence",
    "Description",
    "Note",
    "Photo file",
    "Discord message link",
    "Today's total",
    "Today's cars",
    "Today's motorcycles",
]

# A cell whose text starts with one of these can be mistaken for a formula.
FORMULA_STARTS = ("=", "+", "-", "@", "\t", "\r")


@dataclass
class DayCounts:
    """Vehicles identified today. Starts again from zero when the date changes."""

    day: date
    total: int = 0
    cars: int = 0
    motorcycles: int = 0

    def roll_to(self, today: date) -> None:
        if today != self.day:
            self.day, self.total, self.cars, self.motorcycles = today, 0, 0, 0

    def add(self, vehicle_type: str) -> None:
        """Count one vehicle. vehicle_type is 'car' or 'motorcycle' (any case)."""
        kind = vehicle_type.lower()
        if kind == "car":
            self.cars += 1
        elif kind == "motorcycle":
            self.motorcycles += 1
        else:
            return
        self.total += 1


class ExcelLog:
    def __init__(self, data_dir: Path):
        self.path = data_dir / LOG_NAME
        self.pending_path = data_dir / PENDING_NAME
        self._lock = threading.Lock()

    # ---- today's counts -------------------------------------------------------------

    def rebuild_counts(self, today: date) -> DayCounts:
        """Work out today's counts from the Log sheet and the backup file (used at start-up)."""
        counts = DayCounts(day=today)
        rows: list[dict] = []
        try:
            rows += self._read_log_rows()
        except Exception as problem:
            print(f"WARNING: could not read today's rows from the Excel file ({problem}).")
        try:
            rows += self._read_pending()
        except Exception as problem:
            print(f"WARNING: could not read the backup file ({problem}).")

        for row in rows:
            moment = row.get("Time")
            if isinstance(moment, datetime) and moment.date() == today and row.get("Status") == STATUS_OK:
                counts.add(str(row.get("Vehicle type") or ""))
        return counts

    # ---- adding rows ----------------------------------------------------------------

    async def append_row(self, row: dict) -> bool:
        """Add one row without blocking the Discord connection. Never raises."""
        return await asyncio.to_thread(self.append_row_blocking, row)

    def append_row_blocking(self, row: dict) -> bool:
        """Save the row (and any rows waiting in the backup file, first). True if saved to Excel."""
        with self._lock:
            try:
                waiting = self._read_pending()
            except Exception as problem:
                print(f"WARNING: could not read the backup file ({problem}).")
                waiting = []

            batch = waiting + [row]
            for attempt in range(1, SAVE_TRIES + 1):
                try:
                    self._save_rows(batch)
                except Exception as problem:
                    print(
                        f"WARNING: Excel file is open (or cannot be saved: {type(problem).__name__}), "
                        f"retrying... (try {attempt} of {SAVE_TRIES})"
                    )
                    if attempt < SAVE_TRIES:
                        time.sleep(SECONDS_BETWEEN_TRIES)
                    continue
                self._clear_pending()
                if waiting:
                    print(f"Added {len(waiting)} waiting row(s) from the backup file to Excel.")
                return True

            self._keep_in_backup(row)
            return False

    # ---- reading and writing the workbook -------------------------------------------

    def _read_log_rows(self) -> list[dict]:
        if not self.path.exists():
            return []
        workbook = load_workbook(self.path, read_only=True)
        try:
            if SHEET not in workbook.sheetnames:
                return []
            rows = workbook[SHEET].iter_rows(min_row=2, values_only=True)
            return [dict(zip(COLUMNS, values)) for values in rows]
        finally:
            workbook.close()

    def _save_rows(self, rows: list[dict]) -> None:
        """Add rows at the bottom of the Log sheet, saving through a temporary file."""
        if self.path.exists():
            workbook = load_workbook(self.path)
            sheet = workbook[SHEET] if SHEET in workbook.sheetnames else workbook.create_sheet(SHEET)
        else:
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = SHEET
            sheet.append(COLUMNS)
            for cell in sheet[1]:
                cell.font = Font(bold=True)
            sheet.freeze_panes = "A2"

        for row in rows:
            _write_row(sheet, row)

        self.path.parent.mkdir(parents=True, exist_ok=True)
        temporary = self.path.with_name(self.path.name + ".tmp")
        workbook.save(temporary)
        os.replace(temporary, self.path)  # fails if Excel has the file open: caller retries

    # ---- the backup file ------------------------------------------------------------

    def _read_pending(self) -> list[dict]:
        if not self.pending_path.exists():
            return []
        rows = []
        for line in self.pending_path.read_text(encoding="utf-8").splitlines():
            if line.strip():
                row = json.loads(line)
                if row.get("Time"):
                    row["Time"] = datetime.fromisoformat(row["Time"])
                rows.append(row)
        return rows

    def _keep_in_backup(self, row: dict) -> None:
        try:
            saved = dict(row)
            if isinstance(saved.get("Time"), datetime):
                saved["Time"] = saved["Time"].isoformat()
            with open(self.pending_path, "a", encoding="utf-8") as handle:
                handle.write(json.dumps(saved, ensure_ascii=False) + "\n")
            print("WARNING: the Excel file could not be saved. The row is waiting in the backup file")
            print(f"         {self.pending_path} and will be added the next time a save works.")
        except Exception as problem:
            print(f"ERROR: could not save the row to Excel or to the backup file ({problem}).")

    def _clear_pending(self) -> None:
        try:
            self.pending_path.unlink(missing_ok=True)
        except OSError as problem:
            print(f"WARNING: rows were saved, but the backup file could not be removed ({problem}).")


def _write_row(sheet, row: dict) -> None:
    sheet.append([row.get(column) for column in COLUMNS])
    last = sheet.max_row

    # START GUARD 4: text from other people (names, notes, Claude's answer) that starts like a formula is stored as plain text
    for cell in sheet[last]:
        if isinstance(cell.value, str) and cell.value.startswith(FORMULA_STARTS):
            cell.data_type = "s"  # "s" = text: Excel shows it as typed and never runs it
            cell.quotePrefix = True
    # END GUARD 4

    sheet.cell(row=last, column=1).number_format = TIME_FORMAT
